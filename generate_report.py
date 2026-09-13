# Сборка итогового отчёта: Excel + PDF, а также примеры успешных/ошибочных
# предсказаний. Запуск из корня:  python generate_report.py

import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage,
)

from dataset import build_dataloaders
from models import build_model
from train import load_config


def collect_all_metrics(reports_dir: str = "reports") -> pd.DataFrame:
    rows = []
    for jf in Path(reports_dir).glob("eval_*.json"):
        with open(jf, "r", encoding="utf-8") as f:
            m = json.load(f)
        rows.append({
            "Архитектура": m["model_name"],
            "Accuracy": round(m["accuracy"], 4),
            "Precision (macro)": round(m["precision_macro"], 4),
            "Recall (macro)": round(m["recall_macro"], 4),
            "F1 (macro)": round(m["f1_macro"], 4),
            "F1 (weighted)": round(m["f1_weighted"], 4),
            "Время инференса, мс": round(m["inference_time_ms"], 2),
            "Размер модели, МБ": round(m["model_size_mb"], 2),
        })
    if not rows:
        raise RuntimeError("Нет eval_*.json в reports/. Сначала запустите evaluate.py.")
    return pd.DataFrame(rows).sort_values("Accuracy", ascending=False).reset_index(drop=True)


def save_excel(df: pd.DataFrame, path: str = "reports/comparison_table.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Сравнение архитектур"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    center = Alignment(horizontal="center", vertical="center")

    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for row_idx, row in df.iterrows():
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx + 2, column=col_idx, value=value).alignment = center

    for col in ws.columns:
        max_len = max(len(str(c.value)) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    wb.save(path)
    print(f"✓ Excel: {path}")


def save_pdf(df: pd.DataFrame, path: str = "reports/final_report.pdf"):
    doc = SimpleDocTemplate(path, pagesize=landscape(A4))
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Отчёт: распознавание дорожных знаков (GTSRB)", styles["Title"]))
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(Paragraph(
        "Сравнение 5 архитектур нейронных сетей для классификации 43 классов "
        "дорожных знаков. Метрики получены на тестовой выборке GTSRB.",
        styles["Normal"],
    ))
    elements.append(Spacer(1, 0.5 * cm))

    data = [list(df.columns)] + df.astype(str).values.tolist()
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
    ]))
    elements.append(table)

    best = df.iloc[0]["Архитектура"]
    cm_path = Path("reports") / f"confusion_matrix_{best}.png"
    if cm_path.exists():
        elements.append(Spacer(1, 1 * cm))
        elements.append(Paragraph(f"Confusion Matrix — лучшая модель: {best}", styles["Heading2"]))
        elements.append(RLImage(str(cm_path), width=18 * cm, height=15 * cm))

    doc.build(elements)
    print(f"✓ PDF: {path}")


def save_examples(config_path: str, ckpt_path: str, n_success: int = 3, n_fail: int = 3):
    """Сохраняет примеры успешных и ошибочных предсказаний (не менее 3+3)."""
    config = load_config(config_path)
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")

    _, _, test_loader, class_names = build_dataloaders(config)

    model = build_model(config["model_name"], config["num_classes"], pretrained=False)
    ckpt = torch.load(ckpt_path, map_location=device)
    model.load_state_dict(ckpt["model_state_dict"])
    model.to(device).eval()

    successes, failures = [], []
    with torch.no_grad():
        for imgs, labels in test_loader:
            outputs = model(imgs.to(device))
            probs = torch.softmax(outputs, dim=1)
            preds = outputs.argmax(dim=1).cpu()

            for i in range(imgs.size(0)):
                item = (imgs[i], labels[i].item(), preds[i].item(), probs[i].cpu().numpy())
                if preds[i].item() == labels[i].item() and len(successes) < n_success:
                    successes.append(item)
                elif preds[i].item() != labels[i].item() and len(failures) < n_fail:
                    failures.append(item)

            if len(successes) >= n_success and len(failures) >= n_fail:
                break

    def plot_examples(items, title, save_path):
        fig, axes = plt.subplots(1, len(items), figsize=(4 * len(items), 4))
        if len(items) == 1:
            axes = [axes]
        mean = np.array([0.485, 0.456, 0.406])
        std = np.array([0.229, 0.224, 0.225])
        for ax, (img, true, pred, prob) in zip(axes, items):
            img_np = img.permute(1, 2, 0).numpy() * std + mean
            ax.imshow(np.clip(img_np, 0, 1))
            ax.set_title(f"True: {class_names[true]}\nPred: {class_names[pred]} ({prob[pred]:.2%})",
                         fontsize=10)
            ax.axis("off")
        fig.suptitle(title, fontsize=14)
        plt.tight_layout()
        plt.savefig(save_path, dpi=150)
        plt.close()

    Path("reports").mkdir(exist_ok=True)
    plot_examples(successes, "Успешные предсказания", "reports/examples_success.png")
    plot_examples(failures, "Ошибочные предсказания", "reports/examples_failure.png")
    print("✓ Примеры: reports/examples_success.png, reports/examples_failure.png")


if __name__ == "__main__":
    df = collect_all_metrics("reports")
    print(df.to_string(index=False))
    save_excel(df)
    save_pdf(df)
    save_examples("configs/densenet121.yaml", "models/densenet121_best.pth")
    